"""The tracker workbook end to end (needs openpyxl; skipped without it)."""
import hashlib
import shutil
import unittest
from unittest.mock import patch
import zipfile

from support import DemoWorkspace

try:
    import openpyxl
    from openpyxl.utils import get_column_letter
except ImportError:  # the skill's stdlib tests still run
    openpyxl = None

import layout
import tracker
import tracker_model as model
from applications import Store
from job_intake import JobStore
from job_pipeline import Pipeline
from outreach import OutreachStore, mail_link

ACCOUNT = 'alex.rivera@mail.example'
LONG_JD = '=This is captured source text, not a formula.\n' + ('Full description with ₹ and \U0001F600.\n' * 300)


@unittest.skipIf(openpyxl is None, 'openpyxl is not installed')
class TrackerWorkbookTests(unittest.TestCase):
    def setUp(self):
        self.demo = DemoWorkspace()
        self.ws = self.demo.root
        self.which = patch('tracker.shutil.which', return_value=None)  # use this Python's openpyxl, never uv
        self.which.start()
        self.jobs = JobStore(self.ws, test=True)
        self.jobs.ingest({'jobs': [
            {'company': 'Example Corp', 'role': 'Senior Data Analyst', 'location': 'Toronto, ON',
             'jd_summary': '=HYPERLINK("https://example.test","click")', 'full_jd': LONG_JD,
             'email': 'hiring@example.com', 'posted_date': '2026-10-01',
             'sources': [{'type': 'text', 'label': 'Pasted listing', 'text': LONG_JD, 'access': 'readable'}]},
            {'sources': [{'type': 'linkedin', 'label': 'LinkedIn post', 'access': 'blocked',
                          'url': 'https://www.linkedin.com/posts/example_activity-1234567890123456789'}]}]}, 'codex')
        self.rows = self.jobs.show()['jobs']
        self.app = Store(self.ws, test=True).create('Example Corp', 'Senior Data Analyst', LONG_JD, 'fixture',
                                                    'codex')['application_id']
        self.outreach = OutreachStore(self.ws, test=True)
        self.out = self.outreach.record_sent({
            'company': 'Example Corp', 'recipient': 'hiring@example.com', 'roles': ['Senior Data Analyst'],
            'job_ids': [self.rows[0]['job_id']], 'application_id': self.app, 'outreach_type': 'application',
            'subject': 'Senior Data Analyst application',
            'evidence': {'provider': 'gmail', 'message_id': 'sent-1', 'thread_id': 'thread-a',
                         'sent_at': '2026-09-01T10:00:00-04:00', 'url': mail_link(ACCOUNT, 'sent-1'), 'sender': ACCOUNT,
                         'recipients': ['hiring@example.com'], 'labels': ['SENT'],
                         'capture_sha256': hashlib.sha256(b'sent-1').hexdigest()}}, 'codex')['outreach']['outreach_id']
        self.pipeline = Pipeline(self.ws, test=True)
        self.reply('reply-1', '2026-09-05T10:00:00-04:00')
        self.path = self.ws / layout.sandboxed(layout.TRACKER, True)

    def tearDown(self):
        self.which.stop()
        self.demo.close()

    def reply(self, message_id, occurred_at):
        self.pipeline.reconcile([{'message_id': message_id, 'thread_id': 'thread-a', 'occurred_at': occurred_at,
                                  'sender': 'hiring@example.com', 'recipients': [ACCOUNT],
                                  'subject': 'Re: Senior Data Analyst application',
                                  'body': 'Could you share your notice period?', 'signal': 'hr_question',
                                  'quote': 'Could you share your notice period?'}], 'codex')

    def export(self):
        return tracker.export(self.ws, 'codex', test=True)

    def book(self):
        return openpyxl.load_workbook(str(self.path))

    @staticmethod
    def column(sheet, header):
        return get_column_letter(next(s for s in model.sheet_specs([], '')
                                      if s['name'] == sheet)['headers'].index(header) + 1)

    @staticmethod
    def row_of(ws, key):
        return next(r for r in range(model.FIRST_ROW, ws.max_row + 1) if ws.cell(row=r, column=1).value == key)

    def edit(self, changes):
        """Edit cells like a user would, then save (changes: sheet -> [(key, header, value)])."""
        wb = self.book()
        for name, items in changes.items():
            ws = wb[name]
            for key, header, value in items:
                ws['%s%d' % (self.column(name, header), self.row_of(ws, key))] = value
        wb.save(str(self.path))

    def test_round_trip_layout_and_full_text(self):
        receipt = self.export()
        self.assertTrue(receipt['created'])
        self.assertEqual({k: v['rows'] for k, v in receipt['sheets'].items()},
                         {'Jobs': 2, 'Job Details': receipt['sheets']['Job Details']['rows'], 'Outreach': 1, 'Activity': 1})
        wb = self.book()
        self.assertEqual(wb.sheetnames, ['Jobs', 'Job Details', 'Outreach', 'Activity'])
        jobs = wb['Jobs']
        self.assertEqual(jobs['A2'].value, 'Job listings')
        self.assertIn('America/Toronto', jobs['A3'].value)
        self.assertEqual([c.value for c in jobs[model.HEADER_ROW]][:33], [h for _, h, _ in model.JOB_COLUMNS])
        self.assertEqual(jobs['A5'].fill.fgColor.rgb[-6:], '243447')
        self.assertEqual((jobs['A5'].font.name, jobs['A5'].font.bold, jobs['A6'].font.name), ('Arial', True, 'Arial'))
        self.assertEqual(jobs.freeze_panes, 'D6')
        self.assertEqual(jobs.tables['JobListings'].ref, 'A5:AG7')
        added = jobs['%s%d' % (self.column('Jobs', 'Added on'), self.row_of(jobs, self.rows[0]['job_id']))]
        self.assertEqual(added.number_format, 'yyyy-mm-dd hh:mm')
        posted = jobs['%s%d' % (self.column('Jobs', 'Posted date'), self.row_of(jobs, self.rows[0]['job_id']))]
        self.assertEqual((posted.value.date().isoformat(), posted.number_format), ('2026-10-01', 'yyyy-mm-dd'))
        stage = [dv for dv in jobs.data_validations.dataValidation if dv.errorTitle == 'Stage'][0]
        self.assertEqual(str(stage.sqref), 'AE6:AE1048576')
        self.assertEqual(stage.formula1, '"New,Reviewing,Drafted,Applied,Interviewing,Offer,Rejected,Withdrawn,Closed,Skipped"')
        details = wb['Job Details']
        parts = sorted((r[4].value, r[5].value) for r in details.iter_rows(min_row=6)
                       if r[0].value == self.rows[0]['job_id'] and r[3].value == 'Full JD')
        self.assertEqual(''.join(text for _, text in parts), LONG_JD)
        self.assertEqual(wb['Activity'].cell(row=6, column=4).value, 'HR question')
        outreach = wb['Outreach']
        self.assertEqual(outreach.cell(row=6, column=5).value, 'Sent')
        self.assertEqual(outreach['%s6' % self.column('Outreach', 'Reply action')].value, 'Reply to HR')
        again = self.export()
        self.assertEqual((again['sheets']['Jobs']['added'], again['sheets']['Jobs']['updated_cells']), (0, 0))
        self.assertEqual(again['sheets']['Job Details']['added'], 0)
        self.assertTrue(self.path.with_name('Tracker.xlsx.bak').is_file())
        self.assertFalse((self.ws / layout.TRACKER).exists())

    def test_user_edits_are_saved_back_into_the_records(self):
        self.export()
        job_id = self.rows[0]['job_id']
        self.edit({'Jobs': [(job_id, 'Stage', 'interviewing'), (job_id, 'Notes', 'Panel on Thursday')],
                   'Outreach': [(self.out, 'Follow-up date (editable)', '2026-10-20'),
                                (self.out, 'Next action (editable)', 'Send portfolio link'),
                                (self.out, 'Notes (editable)', 'Recruiter is on leave until Monday')]})
        receipt = self.export()
        self.assertEqual(sorted(p['header'] for p in receipt['pulled_edits']),
                         ['Follow-up date (editable)', 'Next action (editable)', 'Notes', 'Notes (editable)', 'Stage'])
        self.assertEqual(receipt['conflicts'], [])
        row = self.jobs.show(job_id)
        self.assertEqual((row['workflow_status'], row['user_notes']), ('Interviewing', 'Panel on Thursday'))
        self.assertEqual(row['history'][-2]['basis'], 'tracker-edit')
        record = self.outreach.get(self.out)
        self.assertEqual((record['follow_up_override'], record['next_action'], record['notes']),
                         ('2026-10-20', 'Send portfolio link', 'Recruiter is on leave until Monday'))
        jobs = self.book()['Jobs']
        self.assertEqual(jobs['AE%d' % self.row_of(jobs, job_id)].value, 'Interviewing')  # canonical stage name
        quiet = self.export()
        self.assertEqual((quiet['pulled_edits'], quiet['conflicts'], quiet['sheets']['Jobs']['updated_cells']), ([], [], 0))
        self.jobs.set_stage(job_id, 'Offer', 'codex')
        self.export()
        jobs = self.book()['Jobs']
        self.assertEqual(jobs['AE%d' % self.row_of(jobs, job_id)].value, 'Offer')

    def test_both_changed_keeps_the_sheet_and_the_records_and_reports_it(self):
        self.export()
        job_id = self.rows[0]['job_id']
        self.edit({'Jobs': [(job_id, 'Stage', 'Applied'), (job_id, 'Email', 'corrected@example.com')]})
        self.jobs.set_stage(job_id, 'Skipped', 'codex')
        for _ in range(2):  # reported until resolved
            receipt = self.export()
            self.assertEqual([(c['header'], c['sheet_value'], c['json_value']) for c in receipt['conflicts']],
                             [('Stage', 'Applied', 'Skipped')])
            self.assertEqual([k['header'] for k in receipt['kept_edits']], ['Email'])
            self.assertEqual(self.jobs.show(job_id)['workflow_status'], 'Skipped')
            jobs = self.book()['Jobs']
            self.assertEqual(jobs['AE%d' % self.row_of(jobs, job_id)].value, 'Applied')
            self.assertEqual(jobs['E%d' % self.row_of(jobs, job_id)].value, 'corrected@example.com')
        self.edit({'Jobs': [(job_id, 'Stage', 'Skipped')]})
        self.assertEqual(self.export()['conflicts'], [])

    def test_text_is_never_a_formula_and_typed_formulas_are_not_saved(self):
        job_id = self.rows[0]['job_id']
        self.jobs.set_notes(job_id, '=1+1', 'codex')
        self.export()
        jobs = self.book()['Jobs']
        number = self.row_of(jobs, job_id)
        for header in ('JD summary', 'Notes'):
            cell = jobs['%s%d' % (self.column('Jobs', header), number)]
            self.assertEqual(cell.data_type, 's')
        self.assertTrue(jobs['L%d' % number].value.startswith('=HYPERLINK'))
        with zipfile.ZipFile(str(self.path)) as archive:
            sheet_xml = archive.read('xl/worksheets/sheet1.xml').decode('utf-8')
        self.assertNotIn('<f>', sheet_xml)
        self.edit({'Jobs': [(job_id, 'Notes', '=2+2')]})
        receipt = self.export()
        self.assertEqual([r['header'] for r in receipt['rejected_edits']], ['Notes'])
        self.assertIn('formula', receipt['rejected_edits'][0]['reason'])
        self.assertEqual(self.jobs.show(job_id)['user_notes'], '=1+1')

    def test_illegal_characters_are_stripped_and_reported(self):
        self.jobs.ingest({'jobs': [{'company': 'Acme Analytics', 'role': 'Analyst\x07 (contract)',
                                    'sources': [{'type': 'text', 'label': 'Pasted', 'text': 'Analyst\x0b role',
                                                 'access': 'readable'}]}]}, 'codex')
        receipt = self.export()
        flagged = {(c['sheet'], c['header']) for c in receipt['sanitized_cells']}
        self.assertIn(('Jobs', 'Role'), flagged)
        self.assertIn(('Job Details', 'Text'), flagged)
        jobs = self.book()['Jobs']
        self.assertIn('Analyst (contract)', [c.value for c in jobs['C']])
        self.assertEqual(self.export()['sheets']['Jobs']['updated_cells'], 0)

    def test_refuses_parts_it_cannot_keep_and_open_workbooks(self):
        self.export()
        original = self.path.read_bytes()
        owner = self.path.with_name('~$Tracker.xlsx')
        owner.write_bytes(b'')
        with self.assertRaisesRegex(RuntimeError, 'open in a spreadsheet app'):
            self.export()
        owner.unlink()
        self.assertEqual(self.path.read_bytes(), original)
        drawn = self.demo.base / 'drawn.xlsx'
        with zipfile.ZipFile(str(self.path)) as source, zipfile.ZipFile(str(drawn), 'w') as target:
            for item in source.infolist():
                target.writestr(item, source.read(item.filename))
            target.writestr('xl/drawings/drawing1.xml', '<xdr:wsDr/>')
        shutil.copyfile(str(drawn), str(self.path))
        before = self.path.read_bytes()
        with self.assertRaisesRegex(RuntimeError, 'cannot keep'):
            self.export()
        self.assertEqual(self.path.read_bytes(), before)

    def test_user_columns_and_sheets_survive_and_bad_layouts_refuse(self):
        self.export()
        wb = self.book()
        jobs = wb['Jobs']
        jobs['AH5'], jobs['AH6'] = 'My column', 'Personal value'
        extra = wb.create_sheet('My notes')
        extra['A1'], extra['B1'] = 'Keep this', '=1+1'
        number = jobs.max_row + 1
        jobs['A%d' % number], jobs['B%d' % number] = 'JOB-000000000000', 'Old company'
        wb.save(str(self.path))
        self.jobs.ingest({'jobs': [{'company': 'Acme Analytics', 'role': 'Product Analyst', 'full_jd': 'Experiments.',
                                    'sources': [{'type': 'text', 'label': 'Pasted', 'text': 'Experiments.',
                                                 'access': 'readable'}]}]}, 'codex')
        receipt = self.export()
        self.assertEqual((receipt['sheets']['Jobs']['added'], receipt['sheets']['Jobs']['rows']), (1, 4))
        wb = self.book()
        self.assertEqual((wb['Jobs']['AH6'].value, wb['My notes']['B1'].value), ('Personal value', '=1+1'))
        self.assertEqual(wb['Jobs']['B%d' % number].value, 'Old company')  # vanished keys are left alone
        good = self.path.read_bytes()
        wb['Jobs']['B5'] = 'Employer'
        wb.save(str(self.path))
        changed = self.path.read_bytes()
        with self.assertRaisesRegex(RuntimeError, 'should be headed "Company"'):
            self.export()
        self.assertEqual(self.path.read_bytes(), changed)
        self.path.write_bytes(good)
        wb = self.book()
        wb['Jobs']['B40'] = 'A row without an ID'
        wb.save(str(self.path))
        with self.assertRaisesRegex(RuntimeError, 'row 40 has no Job ID'):
            self.export()

    def test_deleted_rows_return_and_activity_only_grows(self):
        self.export()
        wb = self.book()
        jobs = wb['Jobs']
        jobs.delete_rows(self.row_of(jobs, self.rows[1]['job_id']))
        wb.save(str(self.path))
        self.reply('reply-2', '2026-09-07T15:00:00+01:00')
        receipt = self.export()
        self.assertEqual(receipt['restored_rows'], [{'sheet': 'Jobs', 'key': self.rows[1]['job_id']}])
        self.assertEqual(receipt['sheets']['Activity'], {'rows': 2, 'added': 1, 'updated_cells': 0, 'created': False})
        self.assertEqual(self.export()['sheets']['Activity']['added'], 0)


if __name__ == '__main__':
    unittest.main()
