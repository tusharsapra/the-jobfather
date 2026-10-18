# /// script
# requires-python = ">=3.9"
# dependencies = ["openpyxl>=3.1,<4", "tzdata"]
# ///
"""Read and write the tracker workbook (openpyxl), only from a projection JSON made by tracker.py.

    tracker_xlsx.py check
    tracker_xlsx.py read  --workbook Tracker.xlsx --projection projection.json --out snapshot.json
    tracker_xlsx.py write --workbook Tracker.xlsx --projection projection.json

A workbook is refused, and left untouched, when it is open in a spreadsheet app, holds parts this
tool cannot keep (drawings, charts, images, cell notes, pivot tables, slicers, threaded comments,
macros), has changed headers or rows without an ID. Saving goes to a temporary file in the same
folder, which is reopened and checked; the workbook is replaced only if it did not change in the
meantime, and the previous version stays as Tracker.xlsx.bak. Text is always stored as text, so
nothing in a job ad can become a formula.
"""
import argparse
from datetime import date, datetime, time, timedelta
import hashlib
import json
import math
import os
from pathlib import Path
import shutil
import sys
import tempfile
import zipfile

import tracker_model as model

UNSUPPORTED = ('xl/drawings/', 'xl/charts/', 'xl/media/', 'xl/pivotTables/', 'xl/pivotCache/', 'xl/slicers/',
               'xl/slicerCaches/', 'xl/threadedComments/', 'xl/persons/', 'xl/embeddings/', 'xl/ctrlProps/',
               'xl/timelines/', 'xl/timelineCaches/', 'xl/vbaProject.bin')
LAST_ROW = 1048576


class Refused(Exception):
    """The workbook was left untouched; the message says what to do."""


def _openpyxl():
    try:
        import openpyxl
    except ImportError:
        raise Refused('openpyxl is not installed for %s' % sys.executable)
    return openpyxl


def sha256(path):
    digest = hashlib.sha256()
    with open(str(path), 'rb') as stream:
        for block in iter(lambda: stream.read(1 << 20), b''):
            digest.update(block)
    return digest.hexdigest()


def refuse_unsafe(path):
    path = Path(path)
    for owner in (path.with_name('~$' + path.name), path.with_name('.~lock.' + path.name + '#')):
        if owner.exists():
            raise Refused('%s is open in a spreadsheet app (%s exists). Close it and retry; if no app has it '
                          'open, delete %s.' % (path.name, owner.name, owner.name))
    if not path.exists():
        return
    try:
        with zipfile.ZipFile(str(path)) as archive:
            names = archive.namelist()
    except zipfile.BadZipFile:
        raise Refused('%s is not a valid .xlsx workbook; move it aside to export a fresh tracker' % path.name)
    found = sorted({part.rstrip('/') for part in UNSUPPORTED for name in names if name.startswith(part)})
    if found:
        raise Refused('%s contains parts this tool cannot keep (%s): drawings, charts, images, cell notes, pivot '
                      'tables, slicers, threaded comments or macros. Nothing was changed. Move them to another '
                      'workbook, or rename this one so a fresh tracker can be exported.' % (path.name, ', '.join(found)))


def _letter(index):
    from openpyxl.utils import get_column_letter
    return get_column_letter(index)


def _kind(spec, header):
    for kind in ('date', 'datetime', 'integer'):
        if header in spec.get(kind, ()):
            return kind
    return None


def read_cell(cell, kind=None):
    """A workbook cell in projection form; formulas are reported as {'formula': text}."""
    value = cell.value
    if value is None:
        return None
    if cell.data_type == 'f' or type(value).__name__ in ('ArrayFormula', 'DataTableFormula'):
        return {'formula': str(getattr(value, 'text', None) or value)}
    if isinstance(value, datetime):
        if kind == 'date' or (kind is None and value.time() == time(0)):
            return {'date': value.date().isoformat()}
        return {'datetime': model._seconds(value).isoformat()}
    if isinstance(value, date):
        return {'date': value.isoformat()}
    if isinstance(value, (time, timedelta)):
        return str(value)
    return value


def check_headers(ws, spec):
    for index, wanted in enumerate(spec['headers'], 1):
        found = ws.cell(row=model.HEADER_ROW, column=index).value
        if found != wanted:
            raise Refused('%s column %s should be headed "%s" but holds %r. Restore that header, or keep your own '
                          'columns to the right of "%s", then retry.'
                          % (ws.title, _letter(index), wanted, found, spec['headers'][-1]))


def scan(ws, spec):
    """Rows keyed by ID: {key: {header: cell}} and {key: row number}. Refuses blank or repeated IDs."""
    width = len(spec['headers'])
    rows, positions = {}, {}
    for cells in ws.iter_rows(min_row=model.FIRST_ROW, max_col=width):
        if all(c.value is None or (isinstance(c.value, str) and not c.value.strip()) for c in cells):
            continue
        number = cells[0].row
        values = {header: read_cell(cells[i], _kind(spec, header)) for i, header in enumerate(spec['headers'])
                  if i < len(cells)}
        key = model.row_key(spec, values)
        if key is None:
            raise Refused('%s row %d has no %s. Fill it in or delete the row, then retry.'
                          % (ws.title, number, ' / '.join(spec['key'])))
        if key in positions:
            raise Refused('%s rows %d and %d have the same %s; delete the copy, then retry.'
                          % (ws.title, positions[key], number, ' / '.join(spec['key'])))
        positions[key] = number
        rows[key] = values
    return rows, positions


def find_table(ws, spec):
    from openpyxl.utils import range_boundaries
    for name in list(ws.tables):
        table = ws.tables[name]
        min_col, min_row, max_col, _ = range_boundaries(table.ref)
        if min_col == 1 and min_row == model.HEADER_ROW:
            if max_col < len(spec['headers']):
                raise Refused('The %s table must cover columns A to %s; widen it in Excel, then retry.'
                              % (ws.title, _letter(len(spec['headers']))))
            return table
    raise Refused('%s has no table starting at A%d. Undo "Convert to Range", or rename the workbook so a fresh '
                  'tracker can be exported.' % (ws.title, model.HEADER_ROW))


class Styles:
    def __init__(self, projection):
        from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
        self.title = Font(name=model.FONT, size=14, bold=True, color='222222')
        self.note = Font(name=model.FONT, size=10, color='666666')
        self.header = Font(name=model.FONT, size=10, bold=True, color='FFFFFF')
        self.fill = PatternFill('solid', fgColor=model.HEADER_FILL)
        self.header_align = Alignment(wrap_text=True, vertical='center', horizontal='center')
        white = Side(style='thin', color='FFFFFF')
        self.header_border = Border(left=white, right=white)
        self.body = Font(name=model.FONT, size=10, color='222222')
        self.link = Font(name=model.FONT, size=10, color='1155CC', underline='single')
        self.wrap = Alignment(wrap_text=True, vertical='top')
        self.formats = projection['formats']


def _single_link(value):
    return isinstance(value, str) and '\n' not in value and value.lower().startswith(('https://', 'http://'))


def set_value(cell, value, spec, header, styles, report, key):
    """Write one projection value; text is always stored as text (never a formula)."""
    from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
    kind = _kind(spec, header)
    if isinstance(value, dict) and 'date' in value:
        cell.value = date.fromisoformat(value['date'])
        cell.number_format = styles.formats['date']
    elif isinstance(value, dict) and 'datetime' in value:
        cell.value = datetime.fromisoformat(value['datetime'])
        cell.number_format = styles.formats['datetime']
    elif isinstance(value, str):
        text = ILLEGAL_CHARACTERS_RE.sub('', value)
        if text != value:
            report['sanitized_cells'].append({'sheet': spec['name'], 'key': key, 'header': header})
        cell.value = text
        cell.data_type = 's'
        cell.number_format = '@'
    elif isinstance(value, bool):
        cell.value = 'TRUE' if value else 'FALSE'
        cell.data_type = 's'
    elif isinstance(value, (int, float)) and not (isinstance(value, float) and not math.isfinite(value)):
        cell.value = value
        cell.number_format = '0' if kind == 'integer' else 'General'
    else:
        cell.value = None
    if header in spec.get('links', ()):
        if _single_link(cell.value):
            cell.hyperlink = cell.value
            cell.font = styles.link
        elif cell.hyperlink is not None:
            cell.hyperlink = None
            cell.font = styles.body


def _detail_height(text):
    lines = sum(max(1, math.ceil(len(line) / 105)) for line in str(text or '').split('\n'))
    return min(400, max(55, lines * 14 + 12))


def write_row(ws, number, spec, values, styles, report, key):
    for index, header in enumerate(spec['headers'], 1):
        cell = ws.cell(row=number, column=index)
        cell.font = styles.body
        cell.alignment = styles.wrap
        set_value(cell, values.get(header), spec, header, styles, report, key)
    height = spec.get('row_height') or _detail_height(values.get('Text'))
    ws.row_dimensions[number].height = height


def update_row(ws, number, spec, values, keep, styles, report, key):
    changed = 0
    for index, header in enumerate(spec['headers'], 1):
        if header in keep:
            continue
        cell = ws.cell(row=number, column=index)
        wanted = values.get(header)
        if model.normalize(read_cell(cell, _kind(spec, header))) == model.normalize(wanted):
            continue
        set_value(cell, wanted, spec, header, styles, report, key)
        changed += 1
    return changed


def create_sheet(wb, spec, styles):
    from openpyxl.worksheet.table import Table, TableStyleInfo
    ws = wb.create_sheet(spec['name'])
    ws.sheet_view.showGridLines = False
    title = ws.cell(row=model.TITLE_ROW, column=1, value=spec['title'])
    title.font = styles.title
    title.data_type = 's'
    ws.row_dimensions[model.TITLE_ROW].height = 25
    note = ws.cell(row=model.NOTE_ROW, column=1, value=spec['note'])
    note.font = styles.note
    note.data_type = 's'
    for index, (header, width) in enumerate(zip(spec['headers'], spec['widths']), 1):
        cell = ws.cell(row=model.HEADER_ROW, column=index, value=header)
        cell.data_type = 's'
        cell.font, cell.fill = styles.header, styles.fill
        cell.alignment, cell.border = styles.header_align, styles.header_border
        ws.column_dimensions[_letter(index)].width = width
    ws.row_dimensions[model.HEADER_ROW].height = 32
    ws.freeze_panes = model.FREEZE
    taken = {name for sheet in wb.worksheets for name in sheet.tables}
    name, suffix = spec['table'], 2
    while name in taken:
        name, suffix = '%s%d' % (spec['table'], suffix), suffix + 1
    table = Table(displayName=name, ref='A%d:%s%d' % (model.HEADER_ROW, _letter(len(spec['headers'])), model.FIRST_ROW))
    table.tableStyleInfo = TableStyleInfo(name='TableStyleLight1', showRowStripes=True, showColumnStripes=False,
                                          showFirstColumn=False, showLastColumn=False)
    ws.add_table(table)
    return ws


def extend_table(table, last_row):
    from openpyxl.utils import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(table.ref)
    ref = '%s%d:%s%d' % (_letter(min_col), min_row, _letter(max_col), max(max_row, last_row, min_row + 1))
    table.ref = ref
    if table.autoFilter is not None:
        table.autoFilter.ref = ref


def apply_validation(ws, spec, report):
    from openpyxl.worksheet.datavalidation import DataValidation
    for header, values in spec.get('validation', {}).items():
        column = _letter(spec['headers'].index(header) + 1)
        ws.data_validations.dataValidation = [
            dv for dv in ws.data_validations.dataValidation
            if not (dv.errorTitle == header and str(dv.sqref).startswith(column + str(model.FIRST_ROW)))]
        formula = '"%s"' % ','.join(values)
        if any(',' in v or '"' in v for v in values) or len(formula) > 257:
            report['notes'].append('The %s list is too long or has commas, so the %s column has no drop-down'
                                   % (header, ws.title))
            continue
        dv = DataValidation(type='list', formula1=formula, allow_blank=True, showErrorMessage=True,
                            errorTitle=header, error='Choose one of the stages listed in config.json tracker.stages.')
        dv.add('%s%d:%s%d' % (column, model.FIRST_ROW, column, LAST_ROW))
        ws.add_data_validation(dv)


def snapshot(path, projection):
    """The workbook's current rows for the merge (only the kept, non-derived columns)."""
    openpyxl = _openpyxl()
    path = Path(path)
    refuse_unsafe(path)
    before = sha256(path)
    wb = openpyxl.load_workbook(str(path))
    result = {'workbook_sha256': before, 'sheets': {}}
    for spec in projection['sheets']:
        if spec['name'] not in wb.sheetnames:
            result['sheets'][spec['name']] = {'present': False, 'rows': {}}
            continue
        ws = wb[spec['name']]
        check_headers(ws, spec)
        find_table(ws, spec)
        rows, _ = scan(ws, spec)
        kept = [h for h in spec['headers'] if h not in spec['derived']]
        result['sheets'][spec['name']] = {'present': True, 'rows': {
            key: {h: values.get(h) for h in kept} for key, values in rows.items()}}
    if sha256(path) != before:
        raise Refused('%s changed while it was being read; run the export again' % path.name)
    return result


def _replace(source, target):
    for attempt in range(6):
        try:
            os.replace(source, target)
            return
        except PermissionError:
            if os.name != 'nt' or attempt == 5:
                raise
            import time as clock
            clock.sleep(0.05)


def write(path, projection):
    openpyxl = _openpyxl()
    path = Path(path)
    refuse_unsafe(path)
    existed = path.exists()
    original = sha256(path) if existed else None
    expected = projection.get('expected_workbook_sha256')
    if existed != bool(expected) or (existed and original != expected):
        raise Refused('%s changed after it was read; run the export again' % path.name)
    wb = openpyxl.load_workbook(str(path)) if existed else openpyxl.Workbook()
    if not existed:
        wb.remove(wb.active)
        wb.properties.creator = 'The Jobfather'
    styles = Styles(projection)
    report = {'sheets': {}, 'sanitized_cells': [], 'notes': []}
    expected_keys = {}
    for spec in projection['sheets']:
        created = spec['name'] not in wb.sheetnames
        ws = create_sheet(wb, spec, styles) if created else wb[spec['name']]
        if not created:
            check_headers(ws, spec)
        table = find_table(ws, spec)
        _, positions = scan(ws, spec)
        keep = projection.get('keep', {}).get(spec['name'], {})
        added = updated = 0
        next_row = max(positions.values(), default=model.HEADER_ROW) + 1
        for row in spec['rows']:
            values = dict(zip(spec['headers'], row))
            key = model.row_key(spec, values)
            if key in positions:
                if spec['mode'] != 'append':
                    updated += update_row(ws, positions[key], spec, values, keep.get(key, ()), styles, report, key)
                continue
            write_row(ws, next_row, spec, values, styles, report, key)
            positions[key] = next_row
            next_row += 1
            added += 1
        extend_table(table, max(positions.values(), default=model.FIRST_ROW))
        apply_validation(ws, spec, report)
        expected_keys[spec['name']] = set(positions)
        report['sheets'][spec['name']] = {'rows': len(positions), 'added': added, 'updated_cells': updated,
                                          'created': created}
    fd, temp = tempfile.mkstemp(prefix='.' + path.stem + '-', suffix='.xlsx', dir=str(path.parent))
    os.close(fd)
    backup = None
    try:
        wb.save(temp)
        check = openpyxl.load_workbook(temp)
        for spec in projection['sheets']:
            ws = check[spec['name']]
            check_headers(ws, spec)
            rows, _ = scan(ws, spec)
            if set(rows) != expected_keys[spec['name']]:
                raise Refused('The saved %s sheet did not reopen with the expected rows; nothing was replaced'
                              % spec['name'])
        if existed and sha256(path) != original:
            raise Refused('%s changed while exporting; nothing was replaced. Run the export again.' % path.name)
        if existed:
            backup = path.with_name(path.name + '.bak')
            shutil.copy2(str(path), str(backup))
        _replace(temp, str(path))
    finally:
        if os.path.exists(temp):
            os.unlink(temp)
    return dict(report, workbook=str(path), sha256=sha256(path), created=not existed,
                backup=str(backup) if backup else None, saved_reopened=True)


def _atomic_json(path, value):
    path = Path(path)
    fd, temp = tempfile.mkstemp(prefix='.' + path.name + '-', dir=str(path.parent))
    try:
        with os.fdopen(fd, 'wb') as stream:
            stream.write(json.dumps(value, ensure_ascii=False).encode('utf-8'))
            stream.flush()
            os.fsync(stream.fileno())
        os.replace(temp, str(path))
    finally:
        if os.path.exists(temp):
            os.unlink(temp)


def main(argv=None):
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    sub = parser.add_subparsers(dest='command', required=True)
    sub.add_parser('check', help='confirm openpyxl is available')
    for name in ('read', 'write'):
        child = sub.add_parser(name)
        child.add_argument('--workbook', type=Path, required=True)
        child.add_argument('--projection', type=Path, required=True)
        if name == 'read':
            child.add_argument('--out', type=Path, required=True)
    args = parser.parse_args(argv)
    try:
        if args.command == 'check':
            result = {'openpyxl': _openpyxl().__version__, 'python': sys.version.split()[0]}
        else:
            projection = json.loads(args.projection.read_bytes().decode('utf-8'))
            if projection.get('schema_version') != model.SCHEMA_VERSION:
                raise Refused('Unsupported tracker projection')
            if args.command == 'read':
                value = snapshot(args.workbook, projection)
                _atomic_json(args.out, value)
                result = {'workbook_sha256': value['workbook_sha256'],
                          'sheets': {k: {'present': v['present'], 'rows': len(v['rows'])}
                                     for k, v in value['sheets'].items()}}
            else:
                result = write(args.workbook, projection)
    except Refused as error:
        print(str(error), file=sys.stderr)
        return 3
    print(json.dumps(result))  # ASCII only, whatever the console encoding
    return 0


if __name__ == '__main__':
    sys.exit(main())
